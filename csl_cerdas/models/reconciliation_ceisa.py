import base64
import io
from datetime import date, datetime

import openpyxl
import xlsxwriter

from odoo import _, fields, models
from odoo.exceptions import UserError

COLUMNS = ("JENIS DOKUMEN", "NOPEN", "TANGGAL DAFTAR", "SERI BARANG", "URAIAN", "QUANTITY", "SATUAN")
NOTES = (
    "Catatan:",
    "1. Jumlah dokumen yang dianalisis minimal 10% dari populasi dokumen periode tersebut "
    "(minimal 5 dokumen, maksimal 30 dokumen):",
    "2. Dalam hal jumlah dokumen kurang dari 5, jumlah dokumen yang dianalisis adalah 100% dari populasi",
    "3. Analisis meliputi kesesuaian no/tanggal dokumen, uraian, jumlah, dan satuan antara IT Inventory "
    "dengan data CEISA",
    "4. Dalam hal analisis 10% dari populasi ditemukan ketidaksesuaian, maka analisis ditingkatkan menjadi "
    "seluruh jumlah populasi",
    "5. Keterangan diisi apabila ditemukan hasil pemeriksaan yang tidak sesuai untuk menjelaskan tentang "
    "ketidaksesuaian tersebut.",
)


class ReconciliationCeisa(models.Model):
    _name = "reconciliation.ceisa"
    _description = "Rekonsiliasi Ceisa"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _rec_name = "periode"

    periode = fields.Char(string="Periode", tracking=1)
    responsible = fields.Char(string="Penanggung Jawab", tracking=1)
    line_ids = fields.One2many("reconciliation.ceisa.line", "reconciliation_id", string="Rekon Line")

    def action_download_template(self):
        return {
            "type": "ir.actions.act_url",
            "url": "/csl_cerdas/static/xlsx/template_rekonsiliasi_ceisa.xlsx",
            "target": "download",
        }


class ReconciliationCeisaLine(models.Model):
    _name = "reconciliation.ceisa.line"
    _description = "Rekonsiliasi Ceisa Line"

    reconciliation_id = fields.Many2one("reconciliation.ceisa", required=True, ondelete="cascade", index=True)
    ceisa_file = fields.Binary(string="Ceisa 4.0", attachment=True)
    ceisa_filename = fields.Char()
    inventory_file = fields.Binary(string="IT Inventory", attachment=True)
    inventory_filename = fields.Char()
    result_file = fields.Binary(string="Hasil", attachment=True, readonly=True, copy=False)
    result_filename = fields.Char(readonly=True, copy=False)

    def action_generate(self):
        self.ensure_one()
        self.check_access("read")
        self.write({
            "result_file": base64.b64encode(self._generate_result()),
            "result_filename": "hasil_perbandingan_rekonsiliasi.xlsx",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{self._name}/{self.id}/result_file/{self.result_filename}?download=true",
            "target": "download",
        }

    def _generate_result(self):
        self.ensure_one()
        if not (self.ceisa_file and self.inventory_file):
            raise UserError(_("Please upload both the Ceisa 4.0 and IT Inventory files."))
        ceisa = self._read_rows(self.ceisa_file, _("Ceisa 4.0"))
        inventory = self._read_rows(self.inventory_file, _("IT Inventory"))
        return self._build_result(ceisa, inventory)

    def _read_rows(self, data, label):
        """Return {(jenis, nopen, seri): row dict} keeping the file order."""
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data)), read_only=True, data_only=True)
        except Exception as e:
            raise UserError(_("%(file)s is not a valid .xlsx file.", file=label)) from e
        rows = workbook.active.iter_rows(values_only=True)
        header = [str(h or "").strip().upper() for h in next(rows, ())]
        missing = [c for c in COLUMNS if c not in header]
        if missing:
            raise UserError(_("%(file)s is missing column(s): %(cols)s", file=label, cols=", ".join(missing)))
        index = [header.index(c) for c in COLUMNS]
        result = {}
        for values in rows:
            jenis, nopen, tanggal, seri, uraian, qty, satuan = (self._clean(values[i]) for i in index)
            if nopen:
                result[jenis, nopen, seri] = {
                    "jenis": jenis, "nopen": nopen, "tanggal": self._to_date(tanggal),
                    "uraian": uraian, "qty": self._to_number(qty), "satuan": satuan,
                }
        workbook.close()
        return result

    @staticmethod
    def _clean(value):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, str):
            return value.strip()
        return "" if value is None else value

    @staticmethod
    def _to_date(value):
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(str(value), fmt).date().isoformat()
            except ValueError:
                continue
        return str(value)

    @staticmethod
    def _to_number(value):
        try:
            return float(value) if not isinstance(value, (int, float)) else value
        except ValueError:
            return value

    def _build_result(self, ceisa, inventory):
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        sheet = workbook.add_worksheet("Laporan Rekonsiliasi CEISA")
        base = {"font_name": "Calibri", "font_size": 12, "valign": "vcenter"}
        text = workbook.add_format(base)
        signature = workbook.add_format({**base, "font_color": "#858585"})
        head = workbook.add_format({**base, "font_size": 14, "bold": True, "align": "center",
                                    "text_wrap": True, "border": 1})
        cell = workbook.add_format({**base, "align": "center", "border": 1})

        for col, width in zip("ABCDEFGHIJKLMNOP", (8, 16, 21, 14, 21, 14, 46, 46, 21, 26, 14, 14, 21, 14, 31, 23)):
            sheet.set_column(f"{col}:{col}", width)

        jenis = sorted({r["jenis"] for r in [*ceisa.values(), *inventory.values()]}, key=str)
        sheet.merge_range("A1:D1", "Data Rekon IT Inventory dengan Ceisa", text)
        sheet.merge_range("A2:E2", "Jenis Dokumen: %s" % ", ".join(f"BC {j}" for j in jenis), text)
        sheet.merge_range("A3:E3", "Jumlah Dok IT Inventory: %s" % len({r["nopen"] for r in inventory.values()}), text)
        sheet.merge_range("A4:E4", "Jumlah Dok Ceisa: %s" % len({r["nopen"] for r in ceisa.values()}), text)

        sheet.merge_range("A6:A7", "NO. URUT", head)
        sheet.merge_range("B6:B7", "Jenis Dokumen", head)
        sheet.merge_range("C6:D6", "Ceisa", head)
        sheet.merge_range("E6:F6", "IT Inventory", head)
        sheet.merge_range("G6:I6", "Uraian Barang", head)
        sheet.merge_range("J6:L6", "Jumlah", head)
        sheet.merge_range("M6:O6", "Satuan", head)
        sheet.merge_range("P6:P7", "KETERANGAN", head)
        sheet.write_row("C7", ("Nomor", "Tanggal", "Nomor", "Tanggal", "Ceisa", "IT Inventory", "Sesuai/Tidak",
                               "Ceisa", "IT Inventory", "Selisih", "Ceisa", "IT Inventory", "Sesuai/Tidak"), head)

        empty = dict.fromkeys(("jenis", "nopen", "tanggal", "uraian", "qty", "satuan"), "")
        keys = list(inventory) + [k for k in ceisa if k not in inventory]
        for seq, key in enumerate(keys, start=1):
            c, i = ceisa.get(key, empty), inventory.get(key, empty)
            row = seq + 7  # 1-based Excel row
            r = row - 1
            sheet.write_row(r, 0, (seq, c["jenis"] or i["jenis"], c["nopen"], c["tanggal"], i["nopen"],
                                   i["tanggal"], c["uraian"], i["uraian"]), cell)
            sheet.write_formula(r, 8, f'=IF(G{row}=H{row},"Sesuai","Tidak Sesuai")', cell,
                                "Sesuai" if c["uraian"] == i["uraian"] else "Tidak Sesuai")
            sheet.write(r, 9, c["qty"], cell)
            sheet.write(r, 10, i["qty"], cell)
            sheet.write_formula(r, 11, f"=J{row}-K{row}", cell)
            sheet.write_row(r, 12, (c["satuan"], i["satuan"]), cell)
            sheet.write_formula(r, 14, f'=IF(M{row}=N{row},"Sesuai","Tidak Sesuai")', cell,
                                "Sesuai" if c["satuan"] == i["satuan"] else "Tidak Sesuai")
            sheet.write(r, 15, "-", cell)

        last = len(keys) + 7
        for offset, note in enumerate(NOTES, start=last + 1):
            sheet.write(offset, 0, note, text)
        sign_row = last + len(NOTES) + 6
        sheet.write(sign_row, 12, "Ditandatangani secara elektronik", signature)
        sheet.write(sign_row + 1, 12, self.reconciliation_id.responsible or "", text)
        workbook.close()
        return output.getvalue()
